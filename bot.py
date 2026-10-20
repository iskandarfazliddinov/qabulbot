"""
Telegram-bot ro'yxatga olish: Bo'stonliq, Parkent va boshqa tumanlar uchun yangi oqim.
Запуск:  start.bat  (или BOT_TOKEN=... ADMIN_IDS=... python bot.py)
"""
import asyncio
import logging
import os
import re
import sqlite3
import time
from datetime import datetime

from aiogram import Bot, Dispatcher, F, Router
from aiogram.client.default import DefaultBotProperties
from aiogram.filters import Command, CommandStart
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.types import (CallbackQuery, FSInputFile, InlineKeyboardButton,
                           InlineKeyboardMarkup, KeyboardButton, Message,
                           ReplyKeyboardMarkup, ReplyKeyboardRemove)
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# ---------------------------------------------------------------- sozlamalar
# config.txt faylidan o'qish (BOT_TOKEN=..., ADMIN_IDS=..., OQIM=...)
_cfg = os.path.join(os.path.dirname(os.path.abspath(__file__)), "config.txt")
if os.path.exists(_cfg):
    for _line in open(_cfg, encoding="utf-8-sig"):
        if "=" in _line and not _line.strip().startswith("#"):
            _k, _v = _line.split("=", 1)
            os.environ.setdefault(_k.strip(), _v.strip().strip('"'))
BOT_TOKEN = os.environ.get("BOT_TOKEN", "")
if not BOT_TOKEN or "VSTAVTE" in BOT_TOKEN:
    raise SystemExit("XATO: config.txt faylida BOT_TOKEN kiritilmagan!")
_ADMIN_LIST = [int(x) for x in os.environ.get("ADMIN_IDS", "").split(",") if x.strip()]
ADMIN_IDS = set(_ADMIN_LIST)
DB_PATH = os.environ.get("DB_PATH", "oqim.db")
OQIM = os.environ.get("OQIM", "III oqim")          # joriy oqim nomi

DISTRICTS = {
    "bostonliq": {"name": "Bo'stonliq", "start": "10-noyabr 2026", "venue": "Wyndham / Zafaron Premium"},
    "parkent":   {"name": "Parkent",    "start": "17-noyabr 2026", "venue": "Kumushkon Resort"},
    "other":     {"name": "Boshqa tuman", "start": "", "venue": ""},
}
DIRECTIONS = {
    "fo": "Mehmonxona qabulxonasi (Front Office)",
    "hk": "Housekeeping",
    "fb": "Umumiy ovqatlantirish xizmatlarini tashkil qilish",
}
EDUCATION = {
    "9":      "9-sinf",
    "11":     "11-sinf (o'rta)",
    "kollej": "O'rta maxsus (kollej/litsey)",
    "oliy":   "Oliy",
    "talaba": "Hozirda talaba",
}
SHIFTS = {
    "am": "Tushgacha (1-smena)",
    "pm": "Tushdan keyin (2-smena)",
}

WELCOME = (
    "<b>Hurmatli fuqaro!</b>\n\n"
    "O'zbekiston Respublikasi Turizm qo'mitasi tashabbusi bilan Toshkent viloyatining "
    "Bo'stonliq va Parkent tumanlarida mehmonxona va umumiy ovqatlantirish sohasi bo'yicha "
    "<b>BEPUL</b> qisqa muddatli malaka oshirish va qayta tayyorlash kurslari tashkil etilmoqda.\n\n"
    "Kurslarni turizm sohasida kadrlar malakasini oshirish va qayta tayyorlash bo'yicha "
    "<b>yagona akkreditatsiyadan o'tgan muassasa</b> — «Ipak yo'li» turizm va madaniy meros "
    "xalqaro universiteti huzuridagi <b>Xalqaro turizm akademiyasi</b> o'tkazadi. "
    "Kursni muvaffaqiyatli tamomlagan tinglovchilarga sertifikat beriladi.\n\n"
    "<b>Ishtirok etish shartlari:</b>\n"
    "1. Kamida 9-sinfni tamomlagan bo'lish — <b>majburiy</b>.\n"
    "2. Chet tillarini (rus, ingliz) bilish ma'qul ko'riladi.\n"
    "3. Ro'yxatdan o'tgan tinglovchi barcha mashg'ulotlarga <b>qatnashishi shart</b>.\n"
    "4. Pasport ma'lumotlari to'g'ri va to'liq kiritilishi lozim.\n\n"
    f"Joriy oqim: <b>{OQIM}</b>.\nRo'yxatdan o'tish uchun «Boshlash» tugmasini bosing."
)

def cap(s: str) -> str:
    """Har bir so'zning bosh harfi katta (apostrof buzilmaydi): o'rtaboyev -> O'rtaboyev"""
    return " ".join(w[:1].upper() + w[1:].lower() for w in s.split())

PHONE_RE = re.compile(r"^\+?998\d{9}$")
PASSPORT_RE = re.compile(r"^[A-Z]{2}\d{7}$")

logging.basicConfig(level=logging.INFO)
router = Router()

# ---------------------------------------------------------------- ma'lumotlar bazasi
def db():
    con = sqlite3.connect(DB_PATH)
    con.row_factory = sqlite3.Row
    return con

def init_db():
    with db() as con:
        con.execute("""CREATE TABLE IF NOT EXISTS reg (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            created_at TEXT, tg_id INTEGER, tg_username TEXT,
            oqim TEXT, district TEXT, district_name TEXT, direction TEXT,
            fio TEXT, birth TEXT, passport TEXT, education TEXT,
            phone TEXT, phone2 TEXT, shift TEXT, address TEXT, note TEXT)""")
        con.execute("CREATE INDEX IF NOT EXISTS ix_pass ON reg(passport, oqim)")

def save(data: dict, user) -> int:
    with db() as con:
        cur = con.execute(
            """INSERT INTO reg (created_at,tg_id,tg_username,oqim,district,district_name,direction,
               fio,birth,passport,education,phone,phone2,shift,address,note)
               VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (datetime.now().strftime("%Y-%m-%d %H:%M"), user.id, user.username, OQIM,
             data["district"], data["district_name"], data["direction"], data["fio"], data["birth"],
             data["passport"], data["education"], data["phone"], data.get("phone2", ""),
             data["shift"], data["address"], ""))
        return cur.lastrowid

def passport_exists(passport: str) -> bool:
    with db() as con:
        return con.execute("SELECT 1 FROM reg WHERE passport=? AND oqim=?",
                           (passport, OQIM)).fetchone() is not None

# ---------------------------------------------------------------- zaxira nusxa (Telegram orqali)
# Har bir yangi yozuvdan keyin bot bazani admin chatiga yuboradi va "pin" qiladi.
# Server o'chib qolsa, ishga tushganda bot shu pin qilingan fayldan bazani tiklaydi.
BACKUP_CHAT = _ADMIN_LIST[0] if _ADMIN_LIST else None   # birinchi admin — zaxira chati
_dirty = {"since": 0.0}
_last_pin = {"msg_id": None}

def row_count() -> int:
    with db() as con:
        return con.execute("SELECT COUNT(*) FROM reg").fetchone()[0]

async def send_backup(bot: Bot, force: bool = False):
    if not BACKUP_CHAT or not os.path.exists(DB_PATH):
        return
    if not force and not _dirty["since"]:
        return
    try:
        msg = await bot.send_document(
            BACKUP_CHAT, FSInputFile(DB_PATH, filename="oqim.db"),
            caption=f"💾 Zaxira nusxa — {row_count()} ta yozuv, {datetime.now():%d.%m.%Y %H:%M}\n"
                    f"(bu xabarni o'chirmang, bot tiklash uchun ishlatadi)",
            disable_notification=True)
        try:
            if _last_pin["msg_id"]:
                await bot.unpin_chat_message(BACKUP_CHAT, _last_pin["msg_id"])
            await bot.pin_chat_message(BACKUP_CHAT, msg.message_id, disable_notification=True)
            _last_pin["msg_id"] = msg.message_id
        except Exception as e:
            logging.warning("pin failed: %s", e)
        _dirty["since"] = 0.0
    except Exception as e:
        logging.warning("backup failed: %s", e)

async def backup_loop(bot: Bot):
    """Yangi yozuvdan 2 daqiqa o'tgach (bir nechta ro'yxatni birlashtirib) zaxira yuboradi."""
    while True:
        await asyncio.sleep(30)
        if _dirty["since"] and time.time() - _dirty["since"] > 120:
            await send_backup(bot)

async def restore_from_telegram(bot: Bot):
    """Baza bo'sh bo'lsa — admin chatidagi pin qilingan oqim.db faylini yuklab oladi."""
    if not BACKUP_CHAT:
        return
    if os.path.exists(DB_PATH) and row_count() > 0:
        return
    try:
        chat = await bot.get_chat(BACKUP_CHAT)
        pm = chat.pinned_message
        if pm and pm.document and pm.document.file_name == "oqim.db":
            await bot.download(pm.document, destination=DB_PATH)
            _last_pin["msg_id"] = pm.message_id
            logging.info("Baza Telegramdan tiklandi: %s ta yozuv", row_count())
            await bot.send_message(BACKUP_CHAT, f"♻️ Bot qayta ishga tushdi, baza tiklandi: {row_count()} ta yozuv.")
    except Exception as e:
        logging.warning("restore failed: %s", e)

# ---------------------------------------------------------------- holatlar (FSM)
class Reg(StatesGroup):
    welcome = State()
    district = State()
    district_other = State()
    direction = State()
    fio = State()
    birth = State()
    passport = State()
    education = State()
    phone = State()
    phone2 = State()
    shift = State()
    address = State()
    confirm = State()

def ikb(rows):
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text=t, callback_data=c) for t, c in row] for row in rows])

def district_kb():
    return ikb([[(f"📍 {d['name']}", f"dist:{k}")] for k, d in DISTRICTS.items()])

def direction_kb():
    return ikb([[(v, f"dir:{k}")] for k, v in DIRECTIONS.items()])

def education_kb():
    return ikb([[(v, f"edu:{k}")] for k, v in EDUCATION.items()])

def shift_kb():
    return ikb([[(v, f"shift:{k}")] for k, v in SHIFTS.items()])

def phone_kb():
    return ReplyKeyboardMarkup(resize_keyboard=True, one_time_keyboard=True, keyboard=[
        [KeyboardButton(text="📱 Telefon raqamimni yuborish", request_contact=True)]])

def summary(d: dict) -> str:
    return (f"<b>Ma'lumotlaringizni tekshiring:</b>\n\n"
            f"Oqim: {OQIM}\n"
            f"Tuman: {d['district_name']}\n"
            f"Yo'nalish: {DIRECTIONS[d['direction']]}\n"
            f"F.I.O.: {d['fio']}\n"
            f"Tug'ilgan sana: {d['birth']}\n"
            f"Pasport: {d['passport']}\n"
            f"Ma'lumoti: {EDUCATION[d['education']]}\n"
            f"Telefon: {d['phone']}\n"
            f"Qo'shimcha telefon: {d.get('phone2') or '—'}\n"
            f"Smena: {SHIFTS[d['shift']]}\n"
            f"Manzil: {d['address']}")

# ---------------------------------------------------------------- foydalanuvchi oqimi
@router.message(CommandStart())
async def start(m: Message, state: FSMContext):
    await state.clear()
    await m.answer(WELCOME, reply_markup=ikb([[("▶️ Boshlash", "go")]]))
    await state.set_state(Reg.welcome)

@router.message(Command("cancel"))
async def cancel(m: Message, state: FSMContext):
    await state.clear()
    await m.answer("Bekor qilindi. Qaytadan boshlash uchun /start", reply_markup=ReplyKeyboardRemove())

@router.callback_query(Reg.welcome, F.data == "go")
async def go(c: CallbackQuery, state: FSMContext):
    await c.message.answer("Qaysi tumanda o'qiysiz?", reply_markup=district_kb())
    await state.set_state(Reg.district)
    await c.answer()

@router.callback_query(Reg.district, F.data.startswith("dist:"))
async def pick_district(c: CallbackQuery, state: FSMContext):
    key = c.data.split(":")[1]
    await state.update_data(district=key)
    if key == "other":
        await c.message.edit_text("Tumaningiz nomini yozing:\n<i>Masalan: Yangiyo'l</i>")
        await state.set_state(Reg.district_other)
    else:
        await state.update_data(district_name=DISTRICTS[key]["name"])
        await c.message.edit_text(f"Tuman: <b>{DISTRICTS[key]['name']}</b>")
        await c.message.answer("Yo'nalishni tanlang:", reply_markup=direction_kb())
        await state.set_state(Reg.direction)
    await c.answer()

@router.message(Reg.district_other, F.text)
async def district_other(m: Message, state: FSMContext):
    name = " ".join(m.text.split()).strip()
    if len(name) < 3:
        return await m.answer("Tuman nomini to'liq yozing.")
    await state.update_data(district_name=cap(name))
    await m.answer("Yo'nalishni tanlang:", reply_markup=direction_kb())
    await state.set_state(Reg.direction)

@router.callback_query(Reg.direction, F.data.startswith("dir:"))
async def pick_direction(c: CallbackQuery, state: FSMContext):
    key = c.data.split(":")[1]
    await state.update_data(direction=key)
    await c.message.edit_text(f"Yo'nalish: <b>{DIRECTIONS[key]}</b>")
    await c.message.answer("Familiya, ism, otasining ismini kiriting (pasportdagidek):\n"
                           "<i>Masalan: Karimov Aziz Baxtiyorovich</i>")
    await state.set_state(Reg.fio)
    await c.answer()

@router.message(Reg.fio, F.text)
async def get_fio(m: Message, state: FSMContext):
    fio = " ".join(m.text.split())
    if len(fio.split()) < 2 or any(ch.isdigit() for ch in fio):
        return await m.answer("Iltimos, to'liq F.I.O. ni harflar bilan kiriting.")
    await state.update_data(fio=cap(fio))
    await m.answer("Tug'ilgan sanangiz (KK.OO.YYYY):\n<i>Masalan: 15.03.1998</i>")
    await state.set_state(Reg.birth)

@router.message(Reg.birth, F.text)
async def get_birth(m: Message, state: FSMContext):
    try:
        d = datetime.strptime(m.text.strip(), "%d.%m.%Y")
        assert 1940 <= d.year <= 2011
    except Exception:
        return await m.answer("Sana noto'g'ri. Format: KK.OO.YYYY, masalan 15.03.1998")
    await state.update_data(birth=d.strftime("%d.%m.%Y"))
    await m.answer("Pasport seriya va raqami (AA1234567):")
    await state.set_state(Reg.passport)

@router.message(Reg.passport, F.text)
async def get_passport(m: Message, state: FSMContext):
    p = m.text.replace(" ", "").upper()
    if not PASSPORT_RE.match(p):
        return await m.answer("Format: 2 ta harf + 7 ta raqam, masalan AA1234567")
    if passport_exists(p):
        return await m.answer("Bu pasport bilan ushbu oqimga allaqachon ro'yxatdan o'tilgan. "
                              "Xato bo'lsa, administratorga murojaat qiling.")
    await state.update_data(passport=p)
    await m.answer("Ma'lumotingiz:", reply_markup=education_kb())
    await state.set_state(Reg.education)

@router.callback_query(Reg.education, F.data.startswith("edu:"))
async def get_education(c: CallbackQuery, state: FSMContext):
    edu = c.data.split(":")[1]
    await state.update_data(education=edu)
    await c.message.edit_text(f"Ma'lumoti: <b>{EDUCATION[edu]}</b>")
    await c.message.answer("Telefon raqamingizni yuboring — pastdagi tugmani bosing 👇",
                           reply_markup=phone_kb())
    await state.set_state(Reg.phone)
    await c.answer()

@router.message(Reg.phone, F.contact)
async def get_phone_contact(m: Message, state: FSMContext):
    if m.contact.user_id != m.from_user.id:
        return await m.answer("Iltimos, o'zingizning raqamingizni tugma orqali yuboring.")
    p = re.sub(r"[^\d]", "", m.contact.phone_number)
    await state.update_data(phone="+" + p)
    await m.answer("Qo'shimcha telefon raqami (+998XXXXXXXXX). Bo'lmasa «yo'q» deb yozing:",
                   reply_markup=ReplyKeyboardRemove())
    await state.set_state(Reg.phone2)

@router.message(Reg.phone)
async def get_phone_other(m: Message):
    await m.answer("Raqamni qo'lda yozish mumkin emas — pastdagi "
                   "«📱 Telefon raqamimni yuborish» tugmasini bosing.", reply_markup=phone_kb())

@router.message(Reg.phone2, F.text)
async def get_phone2(m: Message, state: FSMContext):
    t = m.text.strip().lower()
    if t in ("yo'q", "yoq", "нет", "-", "—"):
        await state.update_data(phone2="")
    else:
        p = re.sub(r"[^\d]", "", t)
        if len(p) == 9:
            p = "998" + p
        if not PHONE_RE.match(p):
            return await m.answer("Raqam noto'g'ri. Format: +998901234567 yoki «yo'q»")
        await state.update_data(phone2="+" + p)
    await m.answer("Qaysi smenada o'qish siz uchun qulay?", reply_markup=shift_kb())
    await state.set_state(Reg.shift)

@router.callback_query(Reg.shift, F.data.startswith("shift:"))
async def get_shift(c: CallbackQuery, state: FSMContext):
    key = c.data.split(":")[1]
    await state.update_data(shift=key)
    await c.message.edit_text(f"Smena: <b>{SHIFTS[key]}</b>")
    await c.message.answer("Yashash manzilingiz (tuman, mahalla, ko'cha):")
    await state.set_state(Reg.address)
    await c.answer()

@router.message(Reg.address, F.text)
async def get_address(m: Message, state: FSMContext):
    await state.update_data(address=" ".join(m.text.split()))
    data = await state.get_data()
    await m.answer(summary(data), reply_markup=ikb([[("✅ Tasdiqlash", "ok"), ("🔄 Qaytadan", "redo")]]))
    await state.set_state(Reg.confirm)

@router.callback_query(Reg.confirm, F.data == "ok")
async def confirm(c: CallbackQuery, state: FSMContext, bot: Bot):
    data = await state.get_data()
    rid = save(data, c.from_user)
    _dirty["since"] = _dirty["since"] or time.time()
    await state.clear()
    dist = DISTRICTS[data["district"]]
    when = f"Boshlanish: {dist['start']}, {dist['venue']}\n\n" if dist["start"] else "\n"
    await c.message.edit_text(
        f"✅ Ro'yxatga olindingiz! Raqamingiz: <b>#{rid}</b>\n\n"
        f"Tuman: {data['district_name']}\nYo'nalish: {DIRECTIONS[data['direction']]}\n"
        f"Smena: {SHIFTS[data['shift']]}\n{when}"
        f"Mashg'ulot jadvali haqida Telegram orqali xabar beramiz. "
        f"Eslatma: mashg'ulotlarga qatnashish majburiy.")
    for admin in ADMIN_IDS:
        try:
            await bot.send_message(admin, f"🆕 #{rid} {data['district_name']} / {DIRECTIONS[data['direction']]}\n"
                                          f"{data['fio']}, {data['phone']}, {SHIFTS[data['shift']]}")
        except Exception:
            pass
    await c.answer()

@router.callback_query(Reg.confirm, F.data == "redo")
async def redo(c: CallbackQuery, state: FSMContext):
    await state.clear()
    await c.message.edit_text("Qaytadan boshlaymiz. Qaysi tumanda o'qiysiz?", reply_markup=district_kb())
    await state.set_state(Reg.district)
    await c.answer()

# ---------------------------------------------------------------- admin
def is_admin(m: Message) -> bool:
    return m.from_user.id in ADMIN_IDS

@router.message(Command("stats"))
async def stats(m: Message):
    if not is_admin(m):
        return
    with db() as con:
        rows = con.execute("SELECT district, direction, shift, COUNT(*) n FROM reg WHERE oqim=? "
                           "GROUP BY district, direction, shift", (OQIM,)).fetchall()
    out = [f"<b>{OQIM} — statistika</b>"]
    for dk, d in DISTRICTS.items():
        sub = [r for r in rows if r["district"] == dk]
        if not sub:
            continue
        out.append(f"\n📍 <b>{d['name']}</b>: {sum(r['n'] for r in sub)}")
        for dirk, dirn in DIRECTIONS.items():
            am = sum(r["n"] for r in sub if r["direction"] == dirk and r["shift"] == "am")
            pm = sum(r["n"] for r in sub if r["direction"] == dirk and r["shift"] == "pm")
            if am + pm:
                out.append(f"   • {dirn}: {am + pm} (tushgacha {am} / tushdan keyin {pm})")
    await m.answer("\n".join(out))

@router.message(Command("backup"))
async def backup_cmd(m: Message, bot: Bot):
    if not is_admin(m):
        return
    await send_backup(bot, force=True)
    await m.answer("Zaxira nusxa yuborildi va pin qilindi.")

@router.message(Command("export"))
async def export(m: Message):
    if not is_admin(m):
        return
    arg = (m.text.split(maxsplit=1) + [""])[1].strip().lower()
    keys = [arg] if arg in DISTRICTS else list(DISTRICTS)
    sent = False
    for k in keys:
        path = build_xlsx(k)
        if path:
            await m.answer_document(FSInputFile(path), caption=f"{OQIM} — {DISTRICTS[k]['name']}")
            sent = True
    if not sent:
        await m.answer("Hozircha yozuvlar yo'q.")

def build_xlsx(district_key: str):
    name = DISTRICTS[district_key]["name"]
    with db() as con:
        rows = con.execute("SELECT * FROM reg WHERE oqim=? AND district=? "
                           "ORDER BY direction, shift, fio", (OQIM, district_key)).fetchall()
    if not rows:
        return None
    wb = Workbook()
    ws = wb.active
    ws.title = "Baza"
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    bold = Font(name="Times New Roman", bold=True, size=11)
    norm = Font(name="Times New Roman", size=11)
    head_fill = PatternFill("solid", fgColor="D9E1F2")

    headers = ["T/r", "F.I.O.", "Tuman", "Tug'ilgan sana", "Pasport", "Ma'lumoti", "Yo'nalish",
               "Smena", "Telefon", "Qo'shimcha telefon", "Manzil", "Ro'yxat sanasi", "Telegram"]
    last = chr(ord("A") + len(headers) - 1)
    ws.merge_cells(f"A1:{last}1")
    ws["A1"] = f"{OQIM} — {name} — ro'yxatga olingan tinglovchilar"
    ws["A1"].font = Font(name="Times New Roman", bold=True, size=13)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells(f"A2:{last}2")
    ws["A2"] = f"Eksport: {datetime.now():%d.%m.%Y %H:%M}"
    ws["A2"].font = norm
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=4, column=col, value=h)
        c.font, c.fill, c.border = bold, head_fill, border
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for i, r in enumerate(rows, 1):
        vals = [i, r["fio"], r["district_name"], r["birth"], r["passport"], EDUCATION[r["education"]],
                DIRECTIONS[r["direction"]], SHIFTS[r["shift"]], r["phone"], r["phone2"], r["address"],
                r["created_at"], f"@{r['tg_username']}" if r["tg_username"] else str(r["tg_id"])]
        for col, v in enumerate(vals, 1):
            c = ws.cell(row=4 + i, column=col, value=v)
            c.font, c.border = norm, border
            c.alignment = Alignment(vertical="center", wrap_text=True)

    for col, w in zip("ABCDEFGHIJKLM", [5, 32, 14, 14, 12, 22, 34, 22, 16, 18, 34, 16, 16]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"
    path = f"{OQIM.replace(' ', '_')}_{name.replace(chr(39), '').replace(' ', '_')}.xlsx"
    wb.save(path)
    return path

# ---------------------------------------------------------------- ishga tushirish
async def main():
    init_db()
    bot = Bot(BOT_TOKEN, default=DefaultBotProperties(parse_mode="HTML"))
    dp = Dispatcher()
    dp.include_router(router)
    await restore_from_telegram(bot)
    asyncio.create_task(backup_loop(bot))
    await dp.start_polling(bot)

if __name__ == "__main__":
    asyncio.run(main())
